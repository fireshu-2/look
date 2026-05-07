from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn

doc = Document()

# Set default style
style = doc.styles["Normal"]
style.font.name = "Arial"
style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
style.font.size = Pt(12)

# Title
heading = doc.add_heading("YOLOv8 智能视频检测系统 - 代码修改与修复过程全记录", level=0)
heading.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
for run in heading.runs:
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    run.font.bold = True
    run.font.size = Pt(22)

doc.add_paragraph("\n")

def add_section(title, content_list):
    h = doc.add_heading(title, level=1)
    for run in h.runs:
        run.font.name = "Arial"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.size = Pt(16)

    for item in content_list:
        p = doc.add_paragraph()
        p.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY

        # Split item into subtitle and text if it has a colon
        if "：" in item and not item.startswith("代码位置"):
            subtitle, text = item.split("：", 1)
            run_sub = p.add_run(subtitle + "：")
            run_sub.font.bold = True
            run_sub.font.name = "Arial"
            run_sub._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

            run_text = p.add_run(text)
            run_text.font.name = "Arial"
            run_text._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        else:
            run_text = p.add_run(item)
            run_text.font.name = "Arial"
            run_text._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

# Sections Data
sections = [
    ("一、OSD 初始化导致的 Heap Buffer Overflow (堆缓冲区溢出) 修复", [
        "代码位置：src/osd-device.cpp",
        "问题现象：系统在初始化 OSD 引擎并尝试读取颜色查找表（LUT）文件时，发生严重的 Segmentation Fault（段错误），导致程序直接崩溃。",
        "根本原因：底层的 Linux Kernel 驱动严格要求 LUT 颜色表为一个包含 256 种颜色的连续内存块（固定占用 1024 字节：256 * 4 Bytes）。而在原本的代码实现中，系统分配了固定的 1024 字节内存 `m_pcolor_lut = new uint8_t[1024]`，但是在读取文件时，却直接使用了文件的实际大小 `fread(m_pcolor_lut, 1, m_file_size, fp)`。一旦传入的 LUT 文件大小超过 1024 字节，就会发生堆内存越界写入，破坏内存结构。",
        "修复方案：引入边界保护逻辑。使用 `std::min<size_t>(m_file_size, 1024)` 来严格限制最大读取长度。即使文件偏大，也只会读取前 1024 字节，彻底杜绝了内存溢出隐患。"
    ]),
    ("二、摄像头画面全屏绿屏 (Green Screen) 异常修复", [
        "代码位置：src/osd-device.cpp 与 src/pipeline_image.cpp",
        "问题现象：在运行人脸检测 Demo 时，视频流画面被一片纯净的绿色所遮挡，完全无法看到底层的摄像头真实画面。",
        "根本原因：在 YUV 色彩空间中，内存被全 0 填充（Y=0, U=0, V=0）时，其转换为 RGB 呈现出来的就是高饱和度的亮绿色。原本的 OSD 驱动默认开启了所有的 4 个图层（0, 1 为画框层，2, 3 为贴图层）。因为没有向贴图层 2 和 3 写入真实的图像数据，其底层的 DMA 缓冲保持了 0x00 的初始状态，从而形成了遮蔽全屏的“绿幕”。此外，硬件 Pipeline 在修改参数后没有被正确触发更新。",
        "修复方案：在 `osd-device.cpp` 中遍历图层初始化时，加入判断逻辑：仅开启前两个正在使用的图形图层（`osd_enable_layer(..., true)`），并显式关闭未使用的后两个图像图层（`osd_enable_layer(..., false)`）。同时，在 `pipeline_image.cpp` 中补充调用 `UpdateOnlineParam()`，确保底层硬件参数状态机同步刷新。"
    ]),
    ("三、C++ 核心代码编译错误与逻辑还原", [
        "代码位置：demo_face_detection.cpp, src/face_detector.cpp, src/utils.cpp",
        "问题现象：在执行 CMake 与 Make 构建时，终端报出大量编译错误，包括变量未定义、收窄转换警告（Narrowing Conversion）等，导致可执行文件无法生成。",
        "根本原因：(1) `demo_face_detection.cpp` 的 `detector.Predict()` 函数调用传入了一个当前作用域根本不存在的变量 `current_threshold`；(2) `src/face_detector.cpp` 在处理检测框时，遗漏了人脸框形状验证的业务逻辑；(3) `src/utils.cpp` 中存在 `int` 类型向 `float` 进行数学运算时不安全的隐式类型转换警告。",
        "修复方案：将 `demo_face_detection.cpp` 中未定义的 `current_threshold` 替换为全局配置结构体中挂载的正确变量 `g_config.confidence_threshold`；在 `src/face_detector.cpp` 的解析循环中，重新引入 `valid_person_box` 过滤函数的调用，拦截异常检测框；在 `src/utils.cpp` 中，强制加入 `(float)` 类型显式强转，规范 C++ 代码标准，消除编译器的类型收窄报错。"
    ]),
    ("四、Python 高级技术报告与图表自动生成", [
        "代码位置：generate_report.py (新增自动化脚本)",
        "问题现象：用户提出需要输出一份长达 8000 字级别、且包含完整目录和示意图表的 Microsoft Word 格式的深入技术总结报告，以用于汇报或交付。",
        "修复方案：通过引入 `python-docx` 库和 `matplotlib` 绘图库，编写了强大的文档生成引擎。首先利用 Python 在后台静默绘制了系统架构分层图、YOLOv8 算法网络流程图以及 mAP 准确率柱状图，并将图片保存至本地。随后，利用 `docx` 动态创建文档，编排精美的标题、中文字体样式（黑体、宋体等），并将生成的高清技术图片精确嵌入至文档对应的章节中，最终一键输出了高度格式化的 `YOLOv8_Intrusion_Detection_Report.docx`。"
    ]),
    ("五、项目最终总结", [
        "当前状态：",
        "经过上述所有细致入微的 C++ 底层代码修复、内存溢出杜绝以及 Python 工程化报告的编写，目前本项目不仅可以顺利通过 `make -j4` 编译生成稳定可运行的 `ssne_ai_demo` 程序，且完美解决了困扰硬件板卡的绿屏与崩溃顽疾，业务逻辑达到了工控级部署标准。"
    ])
]

for title, content_list in sections:
    add_section(title, content_list)
    doc.add_paragraph("\n")

doc.save("Code_Modification_Process_Report.docx")
print("Code Modification Process Report generated successfully.")
